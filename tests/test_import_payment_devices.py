from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import import_payment_devices as devices  # noqa: E402

TARGET = "M087-500-14-WWA"
OTHER = "ACCESSORY-KIT"


def _write_po(path: Path, rows: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PO"
    ws.append(["banner"])
    ws.append(["NewCo ID", "Projected Opening Date"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()


def _write_orders(path: Path, booked: list[tuple], shipped: list[tuple]) -> None:
    wb = Workbook()
    booked_ws = wb.active
    booked_ws.title = "Booked Orders"
    booked_ws.append(["Daily Dutch Bros Booked Shipped Orders Report"])
    booked_ws.append(
        [
            "End Customer Name",
            "Item Number",
            "Ordered Qty",
            "Requested Date (Approx. Ship Date)",
            "Shipping Date",
        ]
    )
    for row in booked:
        booked_ws.append(list(row))
    shipped_ws = wb.create_sheet("Shipped Orders")
    shipped_ws.append(["metadata"])
    shipped_ws.append(["End Customer Name", "Item Number", "Shipped Qty", "Requested Date", "Shipping Date"])
    for row in shipped:
        shipped_ws.append(list(row))
    wb.save(path)
    wb.close()


def test_shipped_shop_ignores_matching_booked_row():
    events = devices.reconcile_order_lifecycle(
        booked=[
            {
                "id": "AL0107",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2026, 3, 1),
                "shipping_date": None,
            }
        ],
        shipped=[
            {
                "id": "AL0107",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2026, 3, 1),
                "shipping_date": date(2026, 3, 10),
            }
        ],
    )
    assert len(events) == 1
    assert events[0]["id"] == "AL0107"
    assert events[0]["source"] == "shipped"
    assert events[0]["qty"] == 10
    assert events[0]["date"] == date(2026, 3, 10)


def test_firm_burn_excludes_pre_feb_shipment_for_repeat_shop():
    events = devices.reconcile_order_lifecycle(
        booked=[],
        shipped=[
            {
                "id": "AZ0533",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2025, 2, 1),
                "shipping_date": date(2025, 2, 3),
            },
            {
                "id": "AZ0533",
                "item": TARGET,
                "qty": 4,
                "requested_date": date(2026, 3, 1),
                "shipping_date": date(2026, 3, 2),
            },
        ],
    )
    assert [(event["qty"], event["date"]) for event in events] == [
        (4, date(2026, 3, 2)),
    ]
    result = devices.project_lifecycle(events, pending=[], po_shops=[], today=date(2026, 4, 1))
    assert result["firm_end_inventory"] == 5696


def test_pre_contract_shipment_does_not_fulfill_a_future_po():
    events = devices.reconcile_order_lifecycle(
        booked=[],
        shipped=[
            {
                "id": "AZ0703",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2025, 10, 24),
                "shipping_date": date(2025, 10, 24),
            }
        ],
    )
    pending = devices.pending_shops(
        [{"id": "AZ0703", "opening_date": date(2026, 10, 23)}],
        {event["id"] for event in events},
    )
    assert events == []
    assert [shop["id"] for shop in pending] == ["AZ0703"]


def test_booked_only_shop_uses_ordered_qty_on_extract_date():
    events = devices.reconcile_order_lifecycle(
        booked=[
            {
                "id": "CA2909",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2026, 4, 15),
                "shipping_date": None,
            }
        ],
        shipped=[],
        booked_date=date(2026, 4, 10),
    )
    assert events[0]["source"] == "booked"
    assert events[0]["qty"] == 10
    assert events[0]["date"] == date(2026, 4, 10)


def test_report_extract_date_comes_from_filename():
    path = Path("Daily Dutch Bros Booked Shipped Orders Report2026-09-22-06-33-51.xlsx")
    assert devices.report_extract_date(path) == date(2026, 9, 22)


def test_accessory_item_numbers_are_ignored():
    events = devices.reconcile_order_lifecycle(
        booked=[
            {
                "id": "KS0406",
                "item": OTHER,
                "qty": 4,
                "requested_date": date(2026, 5, 1),
                "shipping_date": None,
            }
        ],
        shipped=[
            {
                "id": "KS0406",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2026, 5, 1),
                "shipping_date": date(2026, 5, 8),
            }
        ],
    )
    assert len(events) == 1
    assert events[0]["qty"] == 10
    assert events[0]["source"] == "shipped"


def test_pending_shops_are_po_ids_absent_from_booked_and_shipped():
    po = [
        {"id": "AL0107", "opening_date": date(2026, 6, 1)},
        {"id": "CA2909", "opening_date": date(2026, 7, 1)},
        {"id": "KS0406", "opening_date": date(2026, 8, 1)},
    ]
    pending = devices.pending_shops(po, {"AL0107", "KS0406"})
    assert [shop["id"] for shop in pending] == ["CA2909"]


def test_network_shops_are_not_filtered_by_name():
    events = devices.reconcile_order_lifecycle(
        booked=[],
        shipped=[
            {
                "id": "AZ0533",
                "item": TARGET,
                "qty": 10,
                "requested_date": date(2026, 3, 1),
                "shipping_date": date(2026, 3, 6),
            }
        ],
    )
    pending = devices.pending_shops(
        [{"id": "SC0901", "opening_date": date(2027, 3, 15)}],
        {event["id"] for event in events},
    )
    result = devices.project_lifecycle(events, pending, pending, today=date(2026, 4, 1))
    assert events[0]["id"] == "AZ0533"
    assert pending[0]["id"] == "SC0901"
    assert result["firm_end_inventory"] == 5690


def test_firm_inventory_starts_at_5700_on_feb_1_and_stops_at_cutoff():
    events = [
        {"id": "AL0107", "qty": 10, "date": date(2026, 3, 1), "source": "shipped"},
        {"id": "TN0703", "qty": 10, "date": date(2026, 3, 15), "source": "booked"},
    ]
    result = devices.project_lifecycle(events, pending=[], po_shops=[], today=date(2026, 4, 1))
    assert result["firm"][0] == {"date": "2026-02-01", "inventory": 5700, "series": "firm"}
    assert result["firm"][-1]["date"] == "2026-04-01"
    assert result["firm"][-1]["inventory"] == 5680
    assert result["firm_end_inventory"] == 5680


def test_pending_demand_starts_at_firm_end_and_burns_14_days_before_opening():
    events = [{"id": "AL0107", "qty": 10, "date": date(2026, 3, 1), "source": "shipped"}]
    pending = [{"id": "CA2909", "opening_date": date(2026, 6, 15)}]
    po = pending + [{"id": "AL0107", "opening_date": date(2026, 5, 1)}]
    result = devices.project_lifecycle(events, pending, po, today=date(2026, 4, 1))
    assert result["projected"][0]["date"] == "2026-04-01"
    assert result["projected"][0]["inventory"] == 5690
    staged = next(point for point in result["projected"] if point["date"] == "2026-06-01")
    assert staged["inventory"] == 5680


def test_overdue_pending_shop_does_not_lower_the_dashed_start_point():
    """The dashed line must begin at the exact firm balance, then step down."""
    events = [{"id": "AL0107", "qty": 10, "date": date(2026, 3, 1), "source": "shipped"}]
    pending = [{"id": "IL0802", "opening_date": date(2026, 4, 5)}]
    result = devices.project_lifecycle(events, pending, pending, today=date(2026, 4, 1))
    assert result["firm_end_inventory"] == 5690
    assert result["projected"][0] == {"date": "2026-04-01", "inventory": 5690, "series": "projected"}
    assert result["projected"][1]["inventory"] == 5680


def test_run_rate_starts_after_latest_po_date():
    result = devices.project_lifecycle(
        events=[],
        pending=[],
        po_shops=[{"id": "AL0107", "opening_date": date(2026, 2, 8)}],
        today=date(2026, 2, 1),
    )
    # 5,700 units at 55/week need 726 days after the latest PO date (Feb 8).
    assert result["zero_date"] == date(2028, 2, 4)
    assert result["projected"][-1]["inventory"] == 0


def test_cli_builds_lifecycle_payload(tmp_path: Path):
    raw = tmp_path / "raw"
    raw.mkdir()
    _write_po(
        raw / "PO_Extract_20260916.xlsx",
        [
            ("AL0107", date(2026, 10, 1)),
            ("CA2909", date(2026, 10, 15)),
            ("KS0406", date(2026, 10, 20)),
        ],
    )
    _write_orders(
        raw / "Daily Dutch Bros Booked Shipped Orders Report20260916.xlsx",
        booked=[
            ("AL0107", TARGET, 10, date(2026, 3, 1), None),
            ("AL0107", OTHER, 2, date(2026, 3, 1), None),
            ("KS0406", TARGET, 10, date(2026, 4, 1), None),
        ],
        shipped=[
            ("AL0107", TARGET, 10, date(2026, 3, 1), date(2026, 3, 12)),
        ],
    )
    out = tmp_path / "payment_devices.json"
    assert devices.main(["--raw", str(raw), "--out", str(out), "--as-of", "2026-04-01"]) == 0
    payload = json.loads(out.read_text())
    assert payload["certified"] is True
    assert payload["summary"]["baseline_inventory"] == 5700
    assert payload["summary"]["firm_inventory"] == 5680
    assert payload["summary"]["shipped_shops"] == 1
    assert payload["summary"]["shipped_qty"] == 10
    assert payload["summary"]["booked_qty"] == 10
    assert "Franchise/Boersma" in payload["assumptions"]["scope"]
    assert payload["summary"]["pending_units"] == 10
    assert payload["series"]["firm"][0]["inventory"] == 5700
    assert payload["series"]["projected"][0]["inventory"] == 5680
    assert payload["gap"]["shops"] == 1
    assert payload["gap"]["units"] == 10


def test_missing_sources_do_not_invent_numbers(tmp_path: Path):
    payload = devices.build_payload(tmp_path)
    assert payload["certified"] is False
    assert payload["summary"] is None
    assert payload["series"]["firm"] == []


def test_current_workbooks_count_pre_contract_po_matches_as_pending():
    payload = devices.build_payload(devices.RAW_DIR, today=date(2026, 9, 24))
    if not payload["certified"]:
        return
    assert payload["summary"]["firm_inventory"] == 3760
    assert payload["summary"]["pending_shops"] == 37
    assert payload["summary"]["pending_units"] == 370
    assert payload["gap"]["shops"] == 37
    assert payload["gap"]["units"] == 370
