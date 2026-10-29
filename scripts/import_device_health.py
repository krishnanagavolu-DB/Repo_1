#!/usr/bin/env python3
"""Publish the Device Health workbook summary for the preview terminals card.

The workbook is the counted snapshot. This script does not recount VHQ or
iTrac, and it does not mix those counts into the purchase-order burn down.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "device-health"
PROCESSED = ROOT / "data" / "processed" / "device_health.json"

# Summary-sheet labels. Indented rows keep the leading spaces stripped.
LABELS = {
    "in shops": "in_shops",
    "active": "in_shops_active",
    "inactive / pending": "in_shops_inactive",
    "at stratix": "at_stratix",
    "spare pool": "spare_pool",
    "depot": "depot",
    "awaiting nso / deployment": "awaiting_nso",
    "held for pushed opening (ca4505)": "held_for_opening",
    "other (ca4505)": "held_for_opening",
    "location not confirmed (in vhq, not in a shop, not at stratix)": "location_not_confirmed",
    "total e285 in the system": "total",
}

BUCKETS = ("spare_pool", "depot", "awaiting_nso", "held_for_opening")
HEADER_DATES = re.compile(
    r"VHQ export dated ([A-Za-z]+ \d+, \d{4}).*iTrac \(Stratix\) export dated ([A-Za-z]+ \d+, \d{4}).*built ([A-Za-z]+ \d+, \d{4})"
)


def latest_workbook(directory: Path) -> Path:
    files = sorted(directory.glob("Device_Health_*.xlsx"))
    if not files:
        raise FileNotFoundError(f"no Device_Health workbook in {directory}")
    return files[-1]


def _iso_date(text: str) -> str:
    return datetime.strptime(text, "%b %d, %Y").date().isoformat()


def _count(value, path: Path, label: str) -> int:
    if not isinstance(value, (int, float)):
        raise ValueError(f"{path.name}: {label} is not a count")
    return int(value)


def read_summary(path: Path) -> tuple[dict, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook["Summary"].iter_rows(values_only=True))
    finally:
        workbook.close()

    dates = {"vhq_as_of": "", "itrac_as_of": "", "as_of": ""}
    counts: dict[str, int] = {}
    section = "locations"
    for row in rows:
        label = str(row[0]).strip().lower() if row and row[0] is not None else ""
        if not label:
            continue
        if label.startswith("vhq export dated"):
            match = HEADER_DATES.search(str(row[0]))
            if not match:
                raise ValueError(f"{path.name}: summary header does not name the VHQ, iTrac, and built dates")
            dates["vhq_as_of"] = _iso_date(match.group(1))
            dates["itrac_as_of"] = _iso_date(match.group(2))
            dates["as_of"] = _iso_date(match.group(3))
            continue
        if label.startswith("stratix: spare pool"):
            section = "availability"
            continue
        if "available to order" in label:
            section = "live"
            continue
        key = LABELS.get(label)
        if section == "locations" and key:
            counts[key] = _count(row[1] if len(row) > 1 else None, path, label)
            continue
        if section == "availability" and key in BUCKETS and len(row) > 3:
            counts[f"{key}_available"] = _count(row[2], path, f"{label} available")
            counts[f"{key}_not_available"] = _count(row[3], path, f"{label} not available")
            continue
        if section == "availability" and label == "total at stratix":
            counts["available_to_ship"] = _count(row[2], path, label)
            counts["not_available"] = _count(row[3], path, label)

    required = set(LABELS.values()) | {"available_to_ship", "not_available"}
    missing = sorted(required - set(counts))
    if missing:
        raise ValueError(f"{path.name}: summary is missing {', '.join(missing)}")
    if not all(dates.values()):
        raise ValueError(f"{path.name}: summary header is missing a date")
    parts = counts["in_shops"] + counts["at_stratix"] + counts["location_not_confirmed"]
    if parts != counts["total"]:
        raise ValueError(f"{path.name}: E285 locations {parts} do not add up to {counts['total']}")
    stratix_parts = sum(counts[key] for key in BUCKETS)
    if stratix_parts != counts["at_stratix"]:
        raise ValueError(f"{path.name}: Stratix buckets {stratix_parts} do not add up to {counts['at_stratix']}")
    shop_parts = counts["in_shops_active"] + counts["in_shops_inactive"]
    if shop_parts != counts["in_shops"]:
        raise ValueError(f"{path.name}: shop statuses {shop_parts} do not add up to {counts['in_shops']}")
    if counts["available_to_ship"] + counts["not_available"] != counts["at_stratix"]:
        raise ValueError(f"{path.name}: Stratix availability does not add up to the Stratix total")
    return counts, dates


def build_payload(directory: Path) -> dict:
    path = latest_workbook(directory)
    counts, dates = read_summary(path)
    return {
        **dates,
        "source": path.name,
        "e285": counts,
        "note": (
            "VHQ and iTrac place each E285 once. Available to ship means iTrac status In Inventory. "
            "Not available means In Repair, Picked, or RMA Received. Location not confirmed is not a lost-device count. "
            "These counts are not the purchase order."
        ),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", type=Path, default=RAW)
    parser.add_argument("--out", type=Path, default=PROCESSED)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    payload = build_payload(args.raw)
    rendered = json.dumps(payload, indent=2) + "\n"
    if args.check:
        current = args.out.read_text(encoding="utf-8") if args.out.is_file() else ""
        if current != rendered:
            print("device_health.json does not match the Device Health workbook.", file=sys.stderr)
            return 1
        print(f"Device health matches {payload['source']}.")
        return 0
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(rendered, encoding="utf-8")
    print(f"Wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
